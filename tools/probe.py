"""Temporary: monthly gold price sources (removed after use)."""
import builtins, re, io
from tracker.http import PoliteClient
OUT = []
def say(*a): OUT.append(" ".join(str(x) for x in a))
http = PoliteClient()
def raw(url, n=300):
    try:
        r = http.get(url, retries=0)
        say("RAW", url, r.status_code, r.headers.get("content-type"), len(r.content), "|", r.text[:n].replace("\n", " / ") if "text" in (r.headers.get("content-type") or "") or "json" in (r.headers.get("content-type") or "") or "xml" in (r.headers.get("content-type") or "") else "")
        return r
    except Exception as e:
        say("ERR", url, type(e).__name__, str(e)[:150])
for page in ["https://www.worldbank.org/en/research/commodity-markets"]:
    r = raw(page, 0)
    if r is not None:
        links = sorted(set(re.findall(r'https?://[^"\']+\.xlsx', r.text)))
        say("XLSX", links[:20])
        mon = [l for l in links if "Monthly" in l and "Historical" in l]
        if mon:
            x = raw(mon[0], 0)
            if x is not None:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(x.content), read_only=True, data_only=True)
                say("SHEETS", wb.sheetnames)
                ws = wb["Monthly Prices"]
                rows = list(ws.iter_rows(values_only=True))
                for i, row in enumerate(rows[:8]):
                    say("ROW", i, [c for c in row][:6], "... gold idx?", [j for j, c in enumerate(row) if isinstance(c, str) and "Gold" in c])
                say("LAST", rows[-2][:3], rows[-1][:3])
for u in ["https://api.imf.org/external/sdmx/2.1/data/PCPS/M.W00.PGOLD.USD?startPeriod=2026-01",
          "https://api.imf.org/external/sdmx/3.0/data/dataflow/IMF.RES/PCPS/+/M.W00.PGOLD.USD?c[TIME_PERIOD]=ge:2026-01",
          "https://dataservices.imf.org/REST/SDMX_JSON.svc/CompactData/PCPS/M.W00.PGOLD.USD?startPeriod=2026",
          "https://fred.stlouisfed.org/graph/fredgraph.csv?id=PGOLDUSDM&cosd=2026-01-01"]:
    raw(u, 400)
text = "\n".join(OUT)
for n, i in enumerate(range(0, len(text), 3500)):
    builtins.print(f"::notice title=Probe {n}::" + text[i:i+3500].replace("%", "%25").replace("\r", "").replace("\n", "%0A"))
